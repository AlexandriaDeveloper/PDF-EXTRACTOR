"""Safe, flexible Excel exports for reviewed Oracle payslip data."""
from __future__ import annotations
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from payroll_core import EARNING, DEDUCTION, INFORMATIONAL, export_amount

def safe_text(value):
    s = "" if value is None else str(value)
    # Text is never an Excel formula, even when an attacker can name an item.
    return "'"+s if s.startswith(("=", "+", "-", "@", "\t", "\r")) else s

def create_excel(path, employees, source, *,
                 selected_indexes=None, selected_items=None,
                 include_gross=False, include_deductions=False, include_net=True):
    """Native XLSX with one row/person, genuine numeric amounts and audit evidence."""
    if selected_indexes is None:
        indexes=list(range(len(employees)))
    else:
        indexes=sorted(set(selected_indexes))
    if not indexes or any(i<0 or i>=len(employees) for i in indexes):
        raise ValueError("اختر موظفًا واحدًا على الأقل")
    available={key for e in employees for key in e["items"]}
    keys=list(dict.fromkeys(selected_items or []))
    if any(k not in available for k in keys):raise ValueError("بند غير موجود في ملف المصدر")
    if not (keys or include_gross or include_deductions or include_net):
        raise ValueError("اختر بندًا أو إجماليًا واحدًا على الأقل")
    if len(keys)>500:raise ValueError("الحد الأقصى 500 بند")
    selected=[employees[i] for i in indexes]
    if len(selected)>1048574:raise ValueError("عدد الصفوف يتجاوز حدود Excel")

    wb=Workbook()
    ws=wb.active;ws.title="ملخص الموظفين";ws.sheet_view.rightToLeft=True
    headers=["اسم الموظف","الرقم القومي","رقم الصرف","كود الجهة المؤسسي","كود الموظف"]
    metrics=[]
    if include_gross:metrics.append(("إجمالي الاستحقاقات","printed_gross"))
    if include_deductions:metrics.append(("إجمالي الاستقطاعات","printed_deductions"))
    if include_net:metrics.append(("الصافي الرسمي","printed_net"))
    headers.extend(title for title,_ in metrics)
    headers.extend(f"{kind} | {code} | {label}" for kind,code,label in keys)
    headers.extend(["حالة المراجعة","ملاحظات","صفحات المصدر"])
    ws.append([safe_text(h) for h in headers])
    numeric_indices=set(range(6,6+len(metrics)+len(keys)))
    for e in selected:
        verified=e.get("status")=="Verified"
        values=[safe_text(e.get("name")),safe_text(e.get("national_id")),
                safe_text(e.get("disbursement_number")),safe_text(e.get("institutional_code")),
                safe_text(e.get("code"))]
        values.extend(float(e[field]) if e.get(field) is not None else None for _,field in metrics)
        values.extend(float(val) if (val:=export_amount(e,key,verified)) is not None else None for key in keys)
        values.extend([e.get("status","ReviewRequired")," | ".join(e.get("warnings",[])),
                       ", ".join(str(p) for p in e.get("pages",[]))])
        ws.append(values)
        row=ws.max_row
        for col in numeric_indices:
            ws.cell(row,col).number_format='#,##0.00;[Red](#,##0.00);"-"'
    # Auditable available-value totals; cannot imply verified full reconciliation.
    totals=["مجموع القيم الرقمية المتاحة"]+[None]*(len(headers)-1)
    for c in numeric_indices:
        vals=[ws.cell(r,c).value for r in range(2,ws.max_row+1)]
        totals[c-1]=round(sum(v for v in vals if isinstance(v,(float,int))),2)
    ws.append(totals)

    detail=wb.create_sheet("تفاصيل البنود");detail.sheet_view.rightToLeft=True
    detail.append(["الموظف","الرقم القومي","رقم الصرف","القسم","كود الحساب","البند","القيمة","الصفحة","النص المصدر"])
    count=0
    for e in selected:
        for entry in e["entries"]:
            item_key=(entry["kind"],entry["account_code"],entry["label"])
            if item_key not in keys and not (
                include_gross and entry["kind"]==EARNING or
                include_deductions and entry["kind"]==DEDUCTION):
                continue
            count+=1
            if count>=1048576:raise ValueError("صفوف تفاصيل أكثر من حد Excel")
            detail.append([safe_text(e["name"]),safe_text(e.get("national_id")),
                           safe_text(e.get("disbursement_number")),entry["kind"],
                           safe_text(entry["account_code"]),safe_text(entry["label"]),
                           float(entry["amount"]),entry["page"],safe_text(entry["raw_text"])])
            detail.cell(detail.max_row,7).number_format='#,##0.00'
    notes=wb.create_sheet("ملاحظات ومطابقة");notes.sheet_view.rightToLeft=True
    notes.append(["المجال","القيمة","ملاحظة"])
    notes.append(["ملف المصدر",safe_text(Path(source).name),""])
    notes.append(["عدد الموظفين المختارين",len(selected),""])
    notes.append(["وقت التصدير UTC",datetime.now(timezone.utc).isoformat(),""])
    notes.append(["حصة الحكومة","معلوماتي","لا تضاف إلى مستحقات الموظف النقدية"])
    notes.append(["قيم ناقصة","فارغة","لا تعاملها كأصفار"])
    notes.append(["مجموع متاح","ليس إجماليًا رسميًا معتمدًا","لا تعتمد على مجموع قيم جزئية عندما توجد تحذيرات"])
    for e in selected:
        for issue in e.get("warnings",[]):
            notes.append(["تحذير موظف",safe_text(e["name"]),safe_text(issue)])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2" if sheet != ws else "F2"
        sheet.auto_filter.ref=sheet.dimensions if sheet.max_row>1 else None
        sheet.row_dimensions[1].height=32
        for cell in sheet[1]:
            cell.fill=PatternFill("solid",fgColor="183B56")
            cell.font=Font(bold=True,color="FFFFFF")
            cell.alignment=Alignment(horizontal="right",wrap_text=True)
        for col in sheet.columns:
            letter=get_column_letter(col[0].column)
            sheet.column_dimensions[letter].width= min(50,max(17,max(len(str(c.value or "")) for c in list(col)[:30])+2))
    wb.save(path)
    return {"employees":len(selected),"details":count,"warnings":sum(len(e.get("warnings",[])) for e in selected)}

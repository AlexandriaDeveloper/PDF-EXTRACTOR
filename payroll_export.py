"""Native RTL XLSX export for selected employees and Oracle payroll items.

Amounts are Decimal through the domain layer. Unverified omissions are blank,
never silently invented zeros. Workbook input is a parsed ImportReport.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Iterable
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from payroll_core import ImportReport, PayrollEmployee, PayrollEntry

MAX_ITEMS = 500


@dataclass(frozen=True)
class ExportSelection:
    employee_indexes: tuple[int, ...] | None = None  # None = all employees
    item_keys: tuple[str, ...] = ()
    gross: bool = False
    deductions: bool = False
    net: bool = True


def available_items(report: ImportReport) -> dict[str, PayrollEntry]:
    items: dict[str, PayrollEntry] = {}
    for e in report.employees:
        for item in e.entries:
            if item.kind in ("Earning", "Deduction", "Informational"):
                items.setdefault(item.choice_key, item)
    return dict(sorted(items.items(), key=lambda kv: (kv[1].kind, kv[1].account_code or "", kv[1].label)))


def _validate(report: ImportReport, selection: ExportSelection) -> tuple[list[PayrollEmployee], dict[str, PayrollEntry]]:
    if not report.employees:
        raise ValueError("لا يوجد موظفون مثبتون للتصدير")
    if not selection.item_keys and not (selection.gross or selection.deductions or selection.net):
        raise ValueError("اختر بندًا أو إجماليًا ماليًا واحدًا على الأقل")
    keys = dict.fromkeys(selection.item_keys)
    if len(keys) > MAX_ITEMS:
        raise ValueError("الحد الأقصى لاختيار البنود هو 500 بند")
    choices = available_items(report)
    if any(k not in choices for k in keys):
        raise ValueError("اختيار بند غير موجود في التقرير")
    indexes = list(range(len(report.employees))) if selection.employee_indexes is None else list(dict.fromkeys(selection.employee_indexes))
    if not indexes or any(not isinstance(i, int) or isinstance(i, bool) or i < 0 or i >= len(report.employees) for i in indexes):
        raise ValueError("يجب تحديد موظف واحد على الأقل من الملف")
    return [report.employees[i] for i in indexes], choices


def _text_cell(ws, row: int, col: int, text: object) -> None:
    cell = ws.cell(row, col)
    cell.value = "" if text is None else str(text)
    # openpyxl interprets strings beginning with '=' as Excel formulas.
    # Force a literal string for ALL untrusted PDF-originating text.
    cell.data_type = "s"


def _money_cell(ws, row: int, col: int, value: Decimal | None) -> None:
    cell = ws.cell(row, col)
    if value is not None:
        cell.value = value
        cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
    else:
        cell.value = None


def _header(ws, labels: Iterable[str]) -> None:
    labels = list(labels)
    for i, title in enumerate(labels, 1):
        _text_cell(ws, 1, i, title)
        c = ws.cell(1, i)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B365D")
        c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 40
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "F2" if len(labels) > 5 else "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}1"
    ws.column_dimensions["A"].width = 35
    for i in range(2, len(labels) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24


def _entry_kind_ar(kind: str) -> str:
    return {"Earning": "استحقاق", "Deduction": "استقطاع", "Informational": "معلوماتي"}.get(kind,kind)


def _selected_amount(person: PayrollEmployee, item_key: str, report_verified: bool) -> Decimal | None:
    found = [entry.amount for entry in person.entries if entry.choice_key == item_key]
    if found:
        return sum(found, Decimal(0))
    return Decimal(0) if report_verified and person.verified else None


def make_workbook(report: ImportReport, selection: ExportSelection) -> Workbook:
    people, choices = _validate(report, selection)
    keys = list(dict.fromkeys(selection.item_keys))
    wb = Workbook()
    summary = wb.active
    summary.title = "ملخص الموظفين"
    detail = wb.create_sheet("تفاصيل البنود")
    reviews = wb.create_sheet("معلومات ومراجعة")

    base = ["اسم الموظف", "الرقم القومي", "رقم الصرف", "كود الجهة المؤسسي", "البنك"]
    financial: list[tuple[str,str]] = []
    if selection.gross: financial.append(("gross", "إجمالي الاستحقاقات الرسمي"))
    if selection.deductions: financial.append(("deductions", "إجمالي الاستقطاعات الرسمي"))
    if selection.net: financial.append(("net", "صافي المرتب الرسمي"))
    headings = [f"{_entry_kind_ar(choices[k].kind)} | {choices[k].account_code or 'بلا كود'} | {choices[k].label}" for k in keys]
    _header(summary, base + [label for _,label in financial] + headings +
            ["حالة المراجعة", "ملاحظات الموظف", "الملف", "صفحات المصدر"])

    m_start = len(base)+1
    m_count = len(financial)+len(keys)
    totals: list[Decimal] = [Decimal(0)] * m_count
    known: list[int] = [0] * m_count
    for r, person in enumerate(people, 2):
        for col, value in enumerate((person.name,person.national_id,
                                    person.disbursement_number,person.institutional_entity_code,
                                    person.bank), 1):
            _text_cell(summary,r,col,value)
        amounts = [getattr(person,attr) for attr,_ in financial] + [
            _selected_amount(person,k,report.verified) for k in keys]
        for offset, amount in enumerate(amounts):
            _money_cell(summary,r,m_start+offset,amount)
            if amount is not None:
                totals[offset] += amount
                known[offset] += 1
        after = m_start + m_count
        for j,t in enumerate(("تم التحقق" if report.verified and person.verified else "تحتاج مراجعة",
                               " | ".join(person.warnings),report.filename,
                               f"{person.start_page}-{person.end_page}")):
            _text_cell(summary,r,after+j,t)
        if r % 2 == 0:
            for cell in summary[r]:
                if not cell.fill.patternType:
                    cell.fill = PatternFill("solid",fgColor="F3F7FB")

    total_row = len(people)+2
    _text_cell(summary,total_row,1,"مجموع القيم المتاحة فقط")
    for offset,total in enumerate(totals):
        _money_cell(summary,total_row,m_start+offset,total if known[offset] else None)
    for c in summary[total_row]:
        c.font=Font(bold=True,color="183D56")
        c.fill=PatternFill("solid",fgColor="DDEBF4")
    summary.auto_filter.ref = f"A1:{get_column_letter(len(base)+m_count+4)}{len(people)+1}"

    detail_columns = ["اسم الموظف","الرقم القومي","رقم الصرف","النوع",
                      "كود الحساب","اسم البند","المبلغ","صفحة المصدر","الملف"]
    _header(detail,detail_columns)
    dr=2
    requested=set(keys)
    for e in people:
        for item in e.entries:
            if (item.choice_key not in requested and
                not (selection.gross and item.kind=="Earning") and
                not (selection.deductions and item.kind=="Deduction")):
                continue
            for col,value in enumerate((e.name,e.national_id,e.disbursement_number,
                                        _entry_kind_ar(item.kind),item.account_code,item.label),1):
                _text_cell(detail,dr,col,value)
            _money_cell(detail,dr,7,item.amount)
            _text_cell(detail,dr,8,item.page)
            _text_cell(detail,dr,9,report.filename)
            dr+=1
    if dr>2:
        detail.auto_filter.ref=f"A1:I{dr-1}"

    _header(reviews,["البيان","القيمة","ملاحظة"])
    information = [
        ("الملف المصدر",report.filename,""),
        ("تاريخ التصدير UTC",datetime.now(timezone.utc).isoformat(),""),
        ("حالة التقرير","تم التحقق" if report.verified else "تحتاج مراجعة","لا تُعتمد المبالغ التي لم تثبت صحتها"),
        ("عدد الموظفين المختارين",len(people),""),
        ("عدد البنود المختارة",len(keys),""),
        ("معلوماتي","حصص جهة العمل/الحكومة","لا تضاف إلى استحقاقات الموظف النقدية"),
        ("القيم الفارغة","غير مثبتة","غياب القيمة لا يعني صفرًا"),
        ("مجموع القيم المتاحة","ليس اعتمادًا ماليًا","المجموع يستبعد القيم غير المثبتة"),
        ("الخصوصية","بيانات مالية وشخصية","لا تنشر الملف أو تشاركه دون تصريح"),
    ]
    information.extend(("تحذير التقرير",w,"") for w in report.warnings)
    information.extend(("تحذير موظف",e.name,w) for e in people for w in e.warnings)
    information.extend((f"تغطية العمود {i+1}",f"{known[i]}/{len(people)}","عدد الموظفين ذوي المبالغ المثبتة") for i in range(m_count))
    for row,values in enumerate(information,2):
        for col,value in enumerate(values,1):
            _text_cell(reviews,row,col,value)
    reviews.column_dimensions["A"].width = 35
    reviews.column_dimensions["B"].width = 46
    reviews.column_dimensions["C"].width = 65
    return wb


def save_excel(path: str | Path, report: ImportReport, selection: ExportSelection) -> Path:
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("الملف الناتج يجب أن يكون Excel بصيغة .xlsx")
    workbook=make_workbook(report,selection)
    workbook.save(target)
    return target

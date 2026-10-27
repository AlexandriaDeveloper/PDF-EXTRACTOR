import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from payroll_core import parse_text_pages, normalize_line_item, amount, EARNING, DEDUCTION, INFORMATIONAL
from payroll_excel import create_excel
from openpyxl import load_workbook

SYNTHETIC = """نظام أوراكل للرواتب
اسم الموظف: موظف اختباري محمد علي رقم الصرف: 30200105-00001 الرقم القومى: 00000000000001
نوع الصرف: دفع الكترونى البنك: بنك تجريبي الفرع: المركز
الرقم التامينى: 000001 الدرجة: مدير عام الادارة: كلية اختبارية
الاستحقاقات
21110442 بدل الجامعة 100.00
21110443 مكافأة 50.00
اجمالي الاستحقاقات 150.00
21120202 حكومة إصابة عمل 5.00
الاستقطاعات
41110201 خصم تجريبي 40.00
اجمالي الاستقطاعات 40.00
الصافي: 110.00
"""
SECOND = """نظام أوراكل للرواتب
اسم الموظف: موظف اختباري ثاني رقم الصرف: 30200105-00002 الرقم القومى: 00000000000002
الاستحقاقات
21110442 بدل الجامعة 20.00
اجمالي الاستحقاقات 20.00
الاستقطاعات
41110201 خصم تجريبي 2.00
اجمالي الاستقطاعات 2.00
الصافي: 18.00
"""

class CoreTests(unittest.TestCase):
    def test_separates_header_from_accounts(self):
        records=parse_text_pages([SYNTHETIC,SECOND])
        self.assertEqual(len(records),2)
        first=records[0]
        self.assertEqual(first["institutional_code"],"30200105")
        self.assertEqual(first["disbursement_number"],"30200105-00001")
        self.assertEqual(first["national_id"],"00000000000001")
        self.assertEqual(first["printed_net"],Decimal("110.00"))
        self.assertEqual(first["status"],"Verified")
        self.assertEqual(len(first["items"]),4)
        self.assertEqual(len([k for k in first["items"] if k[0]==INFORMATIONAL]),1)
        self.assertTrue(all(k[1]!="30200105" for k in first["items"]))

    def test_two_accounts_in_one_visual_row_do_not_lose_second_entry(self):
        doc=SYNTHETIC.replace("21110442 بدل الجامعة 100.00\n21110443 مكافأة 50.00",
            "21110442 بدل الجامعة 100.00 21110443 مكافأة 50.00")
        records=parse_text_pages([doc])
        self.assertEqual(len(records),1)
        self.assertEqual(records[0]["status"],"Verified")
        self.assertIn((EARNING,"21110442","بدل الجامعة"),records[0]["items"])
        self.assertIn((EARNING,"21110443","مكافأة"),records[0]["items"])
        self.assertEqual(records[0]["printed_gross"],Decimal("150.00"))

    def test_continuation_and_repeated_items_need_review(self):
        pages=[SYNTHETIC.replace("الصافي: 110.00",""),
               "رقم الصرف: 30200105-00001 الرقم القومى: 00000000000001\n21120202 حكومة إصابة عمل 5.00\nالصافي: 110.00"]
        records=parse_text_pages(pages)
        self.assertEqual(len(records),1)
        self.assertEqual(records[0]["pages"],[1,2])
        self.assertEqual(records[0]["status"],"ReviewRequired")
        self.assertTrue(any("مكرر" in w for w in records[0]["warnings"]))

    def test_invalid_money_and_metadata_never_count(self):
        self.assertIsNone(amount("bad-value"))
        self.assertIsNone(amount("00000000000001"))
        self.assertIsNone(normalize_line_item("30200105-00001 الرقم القومي: 00000000000001",EARNING))
        self.assertIsNone(normalize_line_item("30200105 نوع الصرف: دفع الكتروني",EARNING))
        self.assertIsNone(normalize_line_item("x مجموع 111.00 21110442",EARNING))

    def test_native_xlsx_scopes_employees_and_items(self):
        records=parse_text_pages([SYNTHETIC,SECOND])
        item=next(k for k in records[0]["items"] if k[1]=="21110442")
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"report.xlsx"
            result=create_excel(path,records,"synthetic.pdf",selected_indexes=[0],
                                 selected_items=[item],include_gross=False,include_deductions=False,include_net=True)
            self.assertEqual(result["employees"],1)
            wb=load_workbook(path)
            self.assertEqual(wb.sheetnames,["ملخص الموظفين","تفاصيل البنود","ملاحظات ومطابقة"])
            ws=wb.worksheets[0]
            self.assertTrue(ws.sheet_view.rightToLeft)
            self.assertEqual(ws.max_row,3)
            self.assertEqual(ws["B2"].value,"00000000000001")
            self.assertEqual(ws["F2"].value,110.0)
            self.assertEqual(ws["G2"].value,100.0)
            self.assertEqual(wb.worksheets[1].max_row,2)

    def test_unverified_item_absence_is_blank_not_zero(self):
        records=parse_text_pages([SYNTHETIC, SECOND.replace("الصافي: 18.00","")])
        key=next(k for k in records[0]["items"] if k[1]=="21110443")
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"blank.xlsx"
            create_excel(p,records,"synthetic.pdf",selected_items=[key],include_net=False)
            wb=load_workbook(p)
            self.assertIsNone(wb.worksheets[0]["F3"].value)

    def test_excel_injection_escaped(self):
        raw=SYNTHETIC.replace("موظف اختباري محمد علي","=1+2")
        # Invalid employee name must never be interpreted as formula if included.
        records=parse_text_pages([SYNTHETIC])
        records[0]["name"]="=1+2"
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/"safe.xlsx"
            create_excel(p,records,"synthetic.pdf")
            wb=load_workbook(p)
            self.assertEqual(wb.worksheets[0]["A2"].data_type,"s")
            self.assertTrue(wb.worksheets[0]["A2"].value.startswith("'"))

if __name__=="__main__":unittest.main()
